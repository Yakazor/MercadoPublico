import json
from datetime import date
import sys
from pathlib import Path

from openpyxl import load_workbook
import pytest

from radar.evaluacion import CONDICIONADO, GO, NO_GO, evaluar, puntaje, zona_logistica
from radar.presupuesto import armar_presupuesto

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))


def _lic(**cambios):
    lic = {"codigo": "1-1-LE26", "nombre": "Sillas", "organismo": "Muni", "region": "Región Metropolitana de Santiago",
           "monto_estimado": 2_000_000, "moneda": "CLP", "tipo": "LE", "dias_restantes": 6,
           "fecha_cierre": "2026-10-30 15:00:00", "rubro_nombre": "Mobiliario"}
    return lic | cambios


def _presupuesto(precio=10_000, coincidencia="exacta", cantidad=100, unidad="Unidad", lineas=1, con_precio=None):
    items = [{"codigo": "1-1-LE26", "linea": n, "producto": "Silla", "cantidad": cantidad, "unidad": unidad}
             for n in range(1, lineas + 1)]
    con_precio = lineas if con_precio is None else con_precio
    mercado = [{"linea": n, "precio_unitario_neto": precio, "coincidencia": coincidencia, "verificado": True,
                "tienda": "Tienda"} for n in range(1, con_precio + 1)]
    return armar_presupuesto(items, mercado)


def test_go_con_precios_exactos_margen_y_tiempo():
    ev = evaluar(_lic(), _presupuesto())
    assert ev.decision == GO
    assert ev.costos["directos"] == 1_000_000
    assert ev.margenes["bruto"] > ev.margenes["neto"] > 0
    assert ev.margenes["equilibrio_neto"] == ev.costos["total"]
    assert ev.riesgo == "Bajo" and ev.factibilidad["financiera"] == "Alta"


def test_aproximados_y_precios_faltantes_condicionan():
    ev = evaluar(_lic(), _presupuesto(coincidencia="aproximada"))
    assert ev.decision == CONDICIONADO
    assert any("aproximados" in c for c in ev.condiciones)
    ev = evaluar(_lic(monto_estimado=8_000_000), _presupuesto(lineas=4, con_precio=2))
    assert ev.decision == CONDICIONADO and ev.cobertura == 0.5


def test_no_go_por_cierre_monto_o_sin_precios():
    assert evaluar(_lic(dias_restantes=0), _presupuesto()).decision == NO_GO
    assert evaluar(_lic(monto_estimado=500_000), _presupuesto()).decision == NO_GO     # oferta muy sobre el monto
    assert evaluar(_lic(), _presupuesto(con_precio=0)).decision == NO_GO
    assert evaluar(_lic(solo_regulados=1), _presupuesto()).decision == NO_GO


def test_costos_de_region_garantia_y_orden():
    cerca = evaluar(_lic(), _presupuesto())
    lejos = evaluar(_lic(region="Región de Magallanes y de la Antártica Chilena"), _presupuesto())
    assert lejos.costos["logistica"] == 4 * cerca.costos["logistica"]
    assert zona_logistica("Región de Ñuble") == "lejana" and zona_logistica(None) == "media"
    grande = evaluar(_lic(monto_estimado=100_000_000), _presupuesto())
    assert grande.costos["garantia"] > 0 and cerca.costos["garantia"] == 0
    assert puntaje(cerca) > puntaje(evaluar(_lic(), _presupuesto(coincidencia="aproximada")))


def test_script_genera_excel_informes_correos_y_log(tmp_path, monkeypatch):
    import evaluar_licitaciones as script

    lics = [_lic(), _lic(codigo="2-2-COT26", tipo="Compra Ágil", plazo_entrega_dias=3, llamado=1, nombre="= a lo solicitado"),
            _lic(codigo="3-3-LE26", dias_restantes=2, fecha_cierre="2026-10-22 15:00:00")]   # ya cerró
    items = [{"codigo": "1-1-LE26", "linea": 1, "producto": "Silla", "descripcion": "Silla de oficina",
              "cantidad": 100, "unidad": "Unidad"}]
    mercado = tmp_path / "mercado"
    mercado.mkdir()
    (mercado / "1-1-LE26.json").write_text(json.dumps({"codigo": "1-1-LE26", "items": [{"linea": 1, "opciones": [
        {"tienda": "Tienda X", "url": "https://x.cl/silla", "precio": 11_900, "iva_incluido": True,
         "unidades_equivalentes": 1, "precio_unitario_neto": 10_000, "coincidencia": "exacta", "verificado": True}]}]}))
    (tmp_path / "items.json").write_text(json.dumps(items))
    (tmp_path / "lics.json").write_text(json.dumps(lics))
    entrada = tmp_path / "entrada.json"
    entrada.write_text(json.dumps({"licitaciones_procesadas": "1-1-LE26", "resultado": "revisada"}))
    salida = tmp_path / "eval"
    # Los días al cierre se cuentan desde hoy (no se usa el dias_restantes publicado).
    monkeypatch.setattr(script, "hoy_chile", lambda: date(2026, 10, 24))
    monkeypatch.setattr(sys, "argv", ["x", "--items", str(tmp_path / "items.json"), "--licitaciones",
                                      str(tmp_path / "lics.json"), "--mercado", str(mercado), "--salida", str(salida),
                                      "--registrar", str(entrada)])
    script.main()

    wb = load_workbook(salida / "Licitaciones_evaluadas.xlsx")
    assert wb.sheetnames == ["Licitaciones", "Evaluación", "Claude Log"]
    codigos = [r[0].value for r in wb["Licitaciones"].iter_rows(min_row=2)]
    assert sorted(codigos) == ["1-1-LE26", "2-2-COT26"]           # la cerrada no entra
    # Un nombre que empieza con «=» queda como texto, no como fórmula.
    igual = [c for f in wb["Licitaciones"].iter_rows() for c in f if c.value == "= a lo solicitado"]
    assert igual and all(c.data_type == "s" for c in igual)
    assert wb["Evaluación"]["C2"].value == GO
    assert [r[1].value for r in wb["Claude Log"].iter_rows(min_row=2)] == ["Claude", "Sistema (script)"]
    correo = (salida / "licitaciones" / "1-1-LE26_correos.md").read_text()
    assert "Asunto:** Solicitud de cotización – Sillas" in correo and "Tienda X" in correo and "- Valor neto." in correo
    assert "🟢 GO" in (salida / "licitaciones" / "1-1-LE26_go_no_go.md").read_text()
    assert "Punto de equilibrio" in (salida / "licitaciones" / "1-1-LE26_informe.md").read_text()
    # Segunda corrida el mismo día: reemplaza la entrada automática y conserva la de Claude.
    monkeypatch.setattr(sys, "argv", sys.argv[:-2])
    script.main()
    log = json.loads((salida / "claude_log.json").read_text())
    assert [x["autor"] for x in log] == ["Claude", "Sistema (script)"]


def test_oferta_sobre_el_monto_es_no_go_y_el_iva_sin_credito_es_costo():
    # Oferta 1.000.000 × 1,05 × 1,2 × 1,19 = 1.499.400 con IVA: un poco sobre 1.400.000 ya es NO GO.
    ev = evaluar(_lic(monto_estimado=1_400_000), _presupuesto())
    assert ev.decision == NO_GO and any("supera el monto" in r for r in ev.razones)
    assert evaluar(_lic(monto_estimado=1_600_000), _presupuesto()).decision != NO_GO
    # Sin crédito fiscal el IVA de la compra suma al costo (no a la mercadería) y baja el margen.
    sin_credito = armar_presupuesto([{"linea": 1, "producto": "Silla", "cantidad": 100, "unidad": "Unidad"}],
                                    [{"linea": 1, "precio_unitario_neto": 10_000, "coincidencia": "exacta",
                                      "verificado": True, "tienda": "Tienda"}], iva_es_costo=True)
    ev = evaluar(_lic(monto_estimado=5_000_000), sin_credito)
    assert sin_credito["iva_compras"] == pytest.approx(190_000) and ev.costos["iva_compras"] == pytest.approx(190_000)
    assert ev.costos["directos"] == pytest.approx(1_000_000)
    assert ev.costos["total"] == pytest.approx(ev.costos["directos"] + ev.costos["iva_compras"] + ev.costos["indirectos"])
    assert sin_credito["oferta_neta"] == pytest.approx(1_000_000 * 1.19 * 1.05 * 1.2)
