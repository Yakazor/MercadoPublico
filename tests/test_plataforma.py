import io
import json
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from radar.evaluacion import NO_GO
from radar.orden_proveedor import totales
from radar.plataforma import (CHILE, COTIZACION, ESTIMADA_OC, compra_venta, datos_orden, evaluar_licitacion, filtrar,
                              item_concreto, leer_datos, opciones, orden_xlsx, preparar, ranking, supuestos_config)

AHORA = datetime(2026, 10, 24, 12, 0, tzinfo=CHILE)
RAIZ = Path(__file__).resolve().parents[1]


def _lic(codigo, **cambios):
    lic = {"codigo": codigo, "nombre": "Sillas de oficina", "organismo": "Municipalidad", "tipo": "LE",
           "region": "Región Metropolitana de Santiago", "monto_estimado": 2_000_000, "moneda": "CLP",
           "fecha_cierre": "2026-10-30 15:00:00", "dias_restantes": 7, "rubro_nombre": "Mobiliario",
           "url": f"https://x/{codigo}"}
    return lic | cambios


def _datos():
    return {
        "licitaciones": [
            _lic("1-1-LE26"),
            _lic("2-2-COT26", nombre="Mochilas escolares", tipo="Compra Ágil", monto_estimado=674_000,
                 fecha_cierre="2026-10-26 10:00:00", plazo_entrega_dias=3.0, llamado=1,
                 direccion_entrega="Av. San Gregorio 494"),
            _lic("3-3-LE26", nombre="Ya cerró", fecha_cierre="2026-10-24 09:00:00"),
            _lic("4-4-LE26", nombre="Útiles de aseo", monto_estimado=None, rubro_nombre="Aseo",
                 region="Región de Valparaíso", fecha_cierre="2026-10-25 12:00:00"),
        ],
        "items": [
            {"codigo": "1-1-LE26", "linea": 1, "producto": "Silla", "descripcion": "Silla ergonómica",
             "cantidad": 10, "unidad": "Unidad", "ref_mediana": 60_000},
            {"codigo": "1-1-LE26", "linea": 2, "producto": "Escritorio", "cantidad": 2, "unidad": "Unidad",
             "ref_mediana": 120_000},
            {"codigo": "2-2-COT26", "linea": 1, "producto": "Mochilas", "descripcion": "MOCHILA ESCOLAR 30 L",
             "cantidad": 15, "unidad": "EA", "ref_mediana": 14_072},
        ],
        "mercado": [
            {"codigo": "1-1-LE26", "linea": 1, "tienda": "Tienda A", "producto_mercado": "Silla Pro",
             "precio_unitario_neto": 40_000, "coincidencia": "exacta", "verificado": True},
            {"codigo": "1-1-LE26", "linea": 1, "tienda": "Tienda B", "precio_unitario_neto": 10,
             "coincidencia": "exacta", "descartado": True},
        ],
        "meta": [{"generado": "2026-10-23T12:44:58-03:00"}],
    }


def test_preparar_deja_las_abiertas_con_dias_recalculados():
    prep = preparar(_datos(), AHORA)
    assert [l["codigo"] for l in prep["licitaciones"]] == ["1-1-LE26", "2-2-COT26", "4-4-LE26"]
    assert {l["codigo"]: l["dias_restantes"] for l in prep["licitaciones"]} == \
        {"1-1-LE26": 6, "2-2-COT26": 2, "4-4-LE26": 1}
    assert len(prep["items"]["1-1-LE26"]) == 2
    assert [op["tienda"] for op in prep["mercado"]["1-1-LE26"]] == ["Tienda A"]     # sin la descartada
    assert prep["meta"]["generado"].startswith("2026-10-23")


def test_filtrar_por_texto_tipo_monto_plazo_y_precios():
    lics = preparar(_datos(), AHORA)["licitaciones"]
    codigos = lambda **k: [l["codigo"] for l in filtrar(lics, **k)]  # noqa: E731
    assert codigos() == ["4-4-LE26", "2-2-COT26", "1-1-LE26"]                       # por fecha de cierre
    assert codigos(texto="utiles ASEO") == ["4-4-LE26"]                             # sin tildes ni mayúsculas
    assert codigos(compra_agil=True) == ["2-2-COT26"]
    assert codigos(monto_min=1_000_000) == ["4-4-LE26", "1-1-LE26"]                 # sin monto: no se descarta
    assert codigos(max_dias=2, compra_agil=False) == ["4-4-LE26"]
    assert codigos(con_precios={"1-1-LE26"}) == ["1-1-LE26"]
    assert codigos(plazo_max=2, compra_agil=True) == []
    assert codigos(llamado=1) == ["2-2-COT26"]
    assert codigos(regiones=["Región de Valparaíso"], rubros=["Aseo"]) == ["4-4-LE26"]


def test_cotizacion_manda_sobre_la_tienda_y_estimacion_con_oc():
    prep = preparar(_datos(), AHORA)
    items, mercado = prep["items"]["1-1-LE26"], prep["mercado"]["1-1-LE26"]
    ops = opciones(items, mercado, {1: 30_000})
    assert [(o["linea"], o["tienda"]) for o in ops] == [(1, COTIZACION)]
    ops = opciones(items, mercado, estimar_oc=True)
    estimada = next(o for o in ops if o["linea"] == 2)
    assert estimada["coincidencia"] == ESTIMADA_OC and not estimada["verificado"]
    p, _ = evaluar_licitacion(prep["licitaciones"][0], items, ops)
    assert p["items"][1]["oferta_unitaria_neta"] == pytest.approx(120_000)          # oferta = mediana del Estado


def test_evaluar_compra_agil_con_cotizacion():
    prep = preparar(_datos(), AHORA)
    lic = next(l for l in prep["licitaciones"] if l["codigo"] == "2-2-COT26")
    items = prep["items"]["2-2-COT26"]
    p, ev = evaluar_licitacion(lic, items, opciones(items, [], {1: 25_000}))
    assert p["oferta_total"] == pytest.approx(25_000 * 15 * 1.05 * 1.2 * 1.19)
    assert ev.cobertura == 1 and ev.decision != NO_GO
    p, ev = evaluar_licitacion(lic, items, opciones(items, []))
    assert p["items_sin_precio"] == 1 and ev.decision == NO_GO


def test_ranking_con_precios_cotizaciones_o_estimacion():
    prep = preparar(_datos(), AHORA)
    assert [l["codigo"] for l, _, _ in ranking(prep)] == ["1-1-LE26"]
    con_cotizacion = ranking(prep, cotizaciones={"2-2-COT26": {1: 20_000}})
    assert {l["codigo"] for l, _, _ in con_cotizacion} == {"1-1-LE26", "2-2-COT26"}
    assert len(ranking(prep, estimar_oc=True)) == 2                                 # 4-4 no tiene ítems


def test_supuestos_de_config_y_orden_desde_el_presupuesto():
    s = supuestos_config({"evaluacion": {"valor_hora": 9000, "otro": 1}})
    assert s["valor_hora"] == 9000 and "otro" not in s
    prep = preparar(_datos(), AHORA)
    lic = next(l for l in prep["licitaciones"] if l["codigo"] == "2-2-COT26")
    items = prep["items"]["2-2-COT26"]
    p, _ = evaluar_licitacion(lic, items, opciones(items, [], {1: 21_008.4}))
    datos = datos_orden(lic, p)
    assert datos["items"] == [{"descripcion": "MOCHILA ESCOLAR 30 L", "unidad": "EA", "cantidad": 15,
                               "precio_unitario_neto": 21_008}]
    assert datos["referencia"]["licitacion"] == "2-2-COT26"
    assert "Av. San Gregorio 494" in datos["condiciones"]["lugar_entrega"]
    lic1 = prep["licitaciones"][0]
    items1 = prep["items"]["1-1-LE26"]
    p1, _ = evaluar_licitacion(lic1, items1, opciones(items1, prep["mercado"]["1-1-LE26"]))
    assert [i["descripcion"] for i in datos_orden(lic1, p1, "Tienda A")["items"]] == ["Silla Pro"]
    wb = load_workbook(io.BytesIO(orden_xlsx(datos | {"numero": "OC-2026-009"})))
    celdas = {c.value for fila in wb["Orden de Compra"].iter_rows() for c in fila}
    assert {"OC-2026-009", "MOCHILA ESCOLAR 30 L", 21_008} <= celdas
    assert totales(datos | {"descuento": 8})["total"] == 15 * 21_008 - 8 + round((15 * 21_008 - 8) * 0.19)


def test_leer_datos_desde_carpeta(tmp_path):
    for nombre, valor in _datos().items():
        (tmp_path / f"{nombre}.json").write_text(json.dumps(valor), encoding="utf-8")
    assert leer_datos(str(tmp_path))["licitaciones"][0]["codigo"] == "1-1-LE26"


def _pagina(modulo: str, funcion: str, raiz: str) -> None:
    """Corre una página de la plataforma dentro de AppTest."""
    import importlib
    import sys

    sys.path.insert(0, raiz + "/plataforma")
    getattr(importlib.import_module(modulo), funcion)()


@pytest.mark.parametrize("modulo, funcion", [("inicio", "pagina"), ("explorar", "licitaciones"),
                                             ("explorar", "compras_agiles"), ("evaluar", "pagina"),
                                             ("orden", "pagina"), ("manuales", "pagina"), ("actualizar", "pagina")])
def test_paginas_cargan_sin_errores(tmp_path, monkeypatch, modulo, funcion):
    pytest.importorskip("streamlit")
    import requests
    from streamlit.testing.v1 import AppTest

    # «Actualizar datos» consulta GitHub: sin red en las pruebas.
    monkeypatch.setattr(requests, "get", lambda url, **k: _Respuesta(200, {"workflow_runs": [
        _run(1, "radar-diario.yml", "completed", "success")]}))

    datos = _datos()
    for lic in datos["licitaciones"]:       # que sigan abiertas el día que corra el test
        lic["fecha_cierre"] = "2099-01-01 12:00:00"
    for nombre, valor in datos.items():
        (tmp_path / f"{nombre}.json").write_text(json.dumps(valor), encoding="utf-8")
    monkeypatch.setenv("RADAR_DATOS", str(tmp_path))
    at = AppTest.from_function(_pagina, args=(modulo, funcion, str(RAIZ)), default_timeout=60).run()
    assert not at.exception, [e.value for e in at.exception]
    assert at.title


def test_app_completa_abre_en_inicio(tmp_path, monkeypatch):
    pytest.importorskip("streamlit")
    from streamlit.testing.v1 import AppTest

    for nombre, valor in _datos().items():
        (tmp_path / f"{nombre}.json").write_text(json.dumps(valor), encoding="utf-8")
    monkeypatch.setenv("RADAR_DATOS", str(tmp_path))
    at = AppTest.from_file(str(RAIZ / "plataforma" / "app.py"), default_timeout=60).run()
    assert not at.exception, [e.value for e in at.exception]
    assert at.title[0].value == "Radar Mercado Público"


class _Respuesta:
    def __init__(self, status, datos=None, cabeceras=None):
        self.status_code, self._datos, self.headers, self.text = status, datos or {}, cabeceras or {}, ""

    def json(self):
        return self._datos

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(self.status_code)


def _run(id_, archivo, status, conclusion=None, creada="2026-10-25T00:10:00Z", fin="2026-10-25T00:13:00Z",
         titulo=None):
    return {"id": id_, "path": f".github/workflows/{archivo}", "name": archivo, "status": status,
            "conclusion": conclusion, "event": "workflow_dispatch", "run_started_at": creada,
            "created_at": creada, "updated_at": fin, "html_url": f"https://github.com/x/{id_}",
            "display_title": titulo or archivo}


def test_corridas_traduce_estados_y_errores(monkeypatch):
    from radar import actualizacion as a

    pedidos = []
    runs = [_run(1, "radar-diario.yml", "completed", "success"), _run(2, "compra-agil.yml", "queued"),
            _run(3, "ordenes-historicas.yml", "in_progress"),
            _run(4, "radar-diario.yml", "in_progress", titulo="Publicar datos (solo exportar)")]
    monkeypatch.setattr(a.requests, "get", lambda url, **k: pedidos.append((url, k)) or
                        _Respuesta(200, {"workflow_runs": runs}))
    lista = a.corridas("tok")
    assert [(c["workflow"], c["estado"], c["resultado"]) for c in lista] == [
        ("radar-diario.yml", "completed", "success"), ("compra-agil.yml", "en_cola", None),
        ("ordenes-historicas.yml", "in_progress", None), ("radar-diario.yml", "in_progress", None)]
    assert [c["solo_exportar"] for c in lista] == [False, False, False, True]
    assert lista[3]["nombre"] == "Publicar datos (sin consultar la API)"
    assert lista[0]["nombre"] == "Licitaciones (radar diario)" and lista[0]["fin"].minute == 13
    assert lista[1]["fin"] is None
    assert pedidos[0][1]["params"]["branch"] == "feature/api-mercadopublico"
    assert pedidos[0][1]["headers"]["Authorization"] == "Bearer tok"
    monkeypatch.setattr(a.requests, "get", lambda url, **k: _Respuesta(401))
    with pytest.raises(a.ErrorGitHub, match="token"):
        a.corridas("malo")
    monkeypatch.setattr(a.requests, "get", lambda url, **k: _Respuesta(403, cabeceras={"x-ratelimit-remaining": "0"}))
    with pytest.raises(a.ErrorGitHub, match="limitó"):
        a.corridas()


def test_actualizar_no_pisa_corridas_en_cola(monkeypatch):
    from radar import actualizacion as a

    lanzados = []
    monkeypatch.setattr(a.requests, "post", lambda url, **k: lanzados.append((url.split("/")[-2], k["json"])) or
                        _Respuesta(204))
    libre = [{"workflow": "radar-diario.yml", "estado": "completed"}]
    assert a.actualizar("tok", libre) == [("compra-agil.yml", "lanzado"), ("radar-diario.yml", "lanzado")]
    assert lanzados == [("compra-agil.yml", {"ref": "feature/api-mercadopublico", "inputs": {}}),
                        ("radar-diario.yml", {"ref": "feature/api-mercadopublico", "inputs": {}})]
    lanzados.clear()
    ocupado = [{"workflow": "compra-agil.yml", "estado": "in_progress"},
               {"workflow": "ordenes-historicas.yml", "estado": "en_cola"}]
    resultado = dict(a.actualizar("tok", ocupado))
    assert resultado["compra-agil.yml"] == "ya está corriendo o en cola"
    assert "misma base" in resultado["radar-diario.yml"] and lanzados == []
    # Una corrida «solo exportar» (la lanza Compras Ágiles para publicar) no impide el radar completo.
    lanzados.clear()
    exportando = [{"workflow": "radar-diario.yml", "estado": "in_progress", "solo_exportar": True},
                  {"workflow": "radar-diario.yml", "estado": "en_cola", "solo_exportar": True}]
    assert dict(a.actualizar("tok", exportando))["radar-diario.yml"] == "lanzado"
    with pytest.raises(a.ErrorGitHub):
        a.lanzar("radar-diario.yml", "")
    with pytest.raises(ValueError):
        a.lanzar("otro.yml", "tok")
    monkeypatch.setattr(a.requests, "post", lambda url, **k: _Respuesta(404))
    with pytest.raises(a.ErrorGitHub, match="Read and write"):
        a.lanzar("radar-diario.yml", "tok")


def test_leer_datos_con_marca_salta_la_cache(monkeypatch):
    import radar.plataforma as plataforma

    urls = []
    monkeypatch.setattr(plataforma.requests, "get", lambda url, params=None, **k: urls.append((url, params)) or
                        _Respuesta(200, []))
    plataforma.leer_datos("https://x/datos/", marca="20261025")
    assert urls[0] == ("https://x/datos/licitaciones.json", {"v": "20261025"}) and len(urls) == 4


def _ficha(raiz: str) -> None:
    """Corre la ficha de la primera licitación (ítems y precios de tiendas) dentro de AppTest."""
    import sys

    sys.path.insert(0, raiz + "/plataforma")
    import comun
    import explorar

    comun.conservar_estado("oc_")   # como app.py en cada corrida
    prep = comun.datos()
    explorar.ficha(prep["licitaciones"][0], prep)


def _ficha_sin_precios(raiz: str) -> None:
    import sys

    sys.path.insert(0, raiz + "/plataforma")
    import comun
    import explorar

    prep = comun.datos()
    explorar.ficha(next(l for l in prep["licitaciones"] if l["codigo"] == "2-2-COT26"), prep)


def test_sin_precios_de_tiendas_muestra_el_pedido_para_claude(tmp_path, monkeypatch):
    pytest.importorskip("streamlit")
    from streamlit.testing.v1 import AppTest

    datos = _datos()
    for lic in datos["licitaciones"]:
        lic["fecha_cierre"] = "2099-01-01 12:00:00"
    for nombre, valor in datos.items():
        (tmp_path / f"{nombre}.json").write_text(json.dumps(valor), encoding="utf-8")
    monkeypatch.setenv("RADAR_DATOS", str(tmp_path))
    ficha = AppTest.from_function(_ficha_sin_precios, args=(str(RAIZ),), default_timeout=60).run()
    assert not ficha.exception, [e.value for e in ficha.exception]
    assert [c.value for c in ficha.code] == ["PRECIOS para: 2-2-COT26"]
    evaluacion = AppTest.from_function(_pagina, args=("evaluar", "pagina", str(RAIZ)), default_timeout=60)
    evaluacion.session_state["codigo"] = "2-2-COT26"
    evaluacion.run()
    assert not evaluacion.exception, [e.value for e in evaluacion.exception]
    assert "PRECIOS para: 2-2-COT26" in [c.value for c in evaluacion.code]


def test_ranking_y_ficha_cargan_sin_errores(tmp_path, monkeypatch):
    pytest.importorskip("streamlit")
    from streamlit.testing.v1 import AppTest

    datos = _datos()
    for lic in datos["licitaciones"]:
        lic["fecha_cierre"] = "2099-01-01 12:00:00"
    for nombre, valor in datos.items():
        (tmp_path / f"{nombre}.json").write_text(json.dumps(valor), encoding="utf-8")
    monkeypatch.setenv("RADAR_DATOS", str(tmp_path))
    ranking_ = AppTest.from_function(_pagina, args=("evaluar", "pagina", str(RAIZ)), default_timeout=60)
    ranking_.session_state["eval_vista"] = "Ranking"
    ranking_.run()
    assert not ranking_.exception, [e.value for e in ranking_.exception]
    ficha = AppTest.from_function(_ficha, args=(str(RAIZ),), default_timeout=60).run()
    assert not ficha.exception, [e.value for e in ficha.exception]
    # Segunda corrida (p. ej. al tocar otra cosa con la ficha abierta): conservar_estado("oc_") no puede
    # reasignar la clave de un botón (StreamlitValueAssignmentNotAllowedError, visto en la plataforma).
    ficha.run()
    assert not ficha.exception, [e.value for e in ficha.exception]
    assert ficha.subheader and ficha.subheader[0].value.startswith("1-1-LE26")


def test_llave_de_acceso_deriva_de_la_clave():
    pytest.importorskip("streamlit")
    import sys
    sys.path.insert(0, str(RAIZ / "plataforma"))
    import comun

    llave = comun.token_acceso("ciruela")
    assert llave == comun.token_acceso("ciruela") and len(llave) == 64
    assert llave != comun.token_acceso("otra-clave") and "ciruela" not in llave


def test_acceso_con_clave_recordarme_y_cerrar_sesion(tmp_path, monkeypatch):
    pytest.importorskip("streamlit")
    from streamlit.testing.v1 import AppTest

    for nombre, valor in _datos().items():
        (tmp_path / f"{nombre}.json").write_text(json.dumps(valor), encoding="utf-8")
    monkeypatch.setenv("RADAR_DATOS", str(tmp_path))
    at = AppTest.from_file(str(RAIZ / "plataforma" / "app.py"), default_timeout=60)
    at.secrets["clave"] = "ciruela"
    at.run()
    assert not at.exception, [e.value for e in at.exception]
    assert at.checkbox[0].label.startswith("Recordarme") and at.checkbox[0].value is True
    at.text_input[0].input("mala")
    at.button[0].click().run()
    assert at.error and "incorrecta" in at.error[0].value
    at.text_input[0].input("ciruela")
    at.button[0].click().run()
    assert not at.exception, [e.value for e in at.exception]
    assert at.title[0].value == "Radar Mercado Público" and not at.text_input
    salir = next(b for b in at.button if b.label == "Cerrar sesión")
    salir.click().run()
    assert at.text_input and at.text_input[0].label == "Clave de acceso"


def _item(descripcion, producto="Producto", cantidad=10, unidad="Unidad", linea=1):
    return {"linea": linea, "producto": producto, "descripcion": descripcion, "cantidad": cantidad, "unidad": unidad}


def test_compra_venta_directa_servicios_isp_y_bases():
    assert compra_venta({}, [_item("Mochila escolar 30 L color azul")])[0] is True
    # «servicios médicos» en la categoría ONU de un botiquín no lo vuelve servicio.
    assert compra_venta({}, [_item("Botiquín tipo maletín", producto="Estuches de primeros auxilios para los "
                                   "servicios médicos de urgencias")])[0] is True
    assert compra_venta({}, [_item("Soporte para monitor de 27 pulgadas")])[0] is True
    for texto in ("Estufa a pellet con servicio de instalación", "Pecheras clínicas con bordado del logo",
                  "Licencias Microsoft Windows 11", "Arriendo de baños químicos por 3 meses"):
        estado, motivo = compra_venta({}, [_item("Silla de oficina ergonómica", linea=1), _item(texto, linea=2)])
        assert estado is False and motivo.startswith("Línea 2"), texto
    assert compra_venta({}, [_item("Curso Excel", producto="Formación, capacitación y entrenamiento")])[0] is False
    estado, motivo = compra_venta({"alertas_regulatorias": "Guantes: registro ISP"}, [_item("Guantes de látex M")])
    assert estado is False and "ISP" in motivo
    assert compra_venta({"rubro": "consultoria_gestion_bi"}, [_item("Informe de gestión")])[0] is False
    assert compra_venta({}, [_item("Insumos de aseo según bases", unidad="Global")])[0] is None
    assert compra_venta({}, [_item("Estufa a pellet modelo X", unidad="GL")])[0] is None    # GL = global
    assert compra_venta({}, [])[0] is None
    assert item_concreto(_item("Soporte para monitor de 27 pulgadas")) and not item_concreto(_item("Soporte técnico anual"))


def test_orden_al_proveedor_sin_precios_estimados_con_oc():
    # Una línea estimada con OC del Estado no tiene precio cotizado: va vacía en la orden.
    prep = preparar(_datos(), AHORA)
    lic = next(l for l in prep["licitaciones"] if l["codigo"] == "2-2-COT26")
    items = [it | {"ref_mediana": 30_000} for it in prep["items"]["2-2-COT26"]]
    p, _ = evaluar_licitacion(lic, items, opciones(items, [], estimar_oc=True))
    assert p["items"][0]["costo_unitario_neto"] > 0
    assert [i["precio_unitario_neto"] for i in datos_orden(lic, p)["items"]] == [None]


@pytest.mark.parametrize("unidad, cotizable", [("Unidad", True), ("Caja", True), ("Kit", False), ("GL", False),
                                               ("Lote", False), ("Unidad no definida", False), ("Servicio", False)])
def test_item_concreto_usa_las_mismas_unidades_que_el_presupuesto(unidad, cotizable):
    item = {"cantidad": 10, "unidad": unidad, "producto": "Silla", "descripcion": "Silla de visita apilable"}
    assert item_concreto(item) is cotizable


def test_dias_al_cierre_desde_hoy():
    from datetime import date
    from radar.plataforma import dias_al_cierre
    hoy = date(2026, 10, 24)
    assert dias_al_cierre({"fecha_cierre": "2026-10-25 15:00:00", "dias_restantes": 3}, hoy) == 1
    assert dias_al_cierre({"fecha_cierre": "2026-10-23", "dias_restantes": 1}, hoy) == -1
    assert dias_al_cierre({"fecha_cierre": None, "dias_restantes": 4}, hoy) == 4


def test_version_datos_lee_el_commit_de_la_rama_y_fija_la_url(monkeypatch):
    import radar.plataforma as plataforma

    sha = "3dac93a05184298919990563edf112fe05105bfd"
    refs = f"001e# service=git-upload-pack\n0000{'a' * 40} HEAD\x00multi_ack\n003e{sha} refs/heads/datos\n"
    pedidas = []

    class Refs:
        text = refs

        def raise_for_status(self):
            pass

    monkeypatch.setattr(plataforma.requests, "get", lambda url, **k: pedidas.append(url) or Refs())
    assert plataforma.version_datos() == sha
    assert pedidas == ["https://github.com/Yakazor/MercadoPublico.git/info/refs"]
    assert plataforma.origen_en_commit(plataforma.ORIGEN_DATOS, sha) == \
        f"https://raw.githubusercontent.com/Yakazor/MercadoPublico/{sha}"
    # Otro origen (carpeta local) o sin commit: queda igual.
    assert plataforma.version_datos("/tmp/datos") is None
    assert plataforma.origen_en_commit("/tmp/datos", sha) == "/tmp/datos"
    assert plataforma.origen_en_commit(plataforma.ORIGEN_DATOS, None) == plataforma.ORIGEN_DATOS

    def caido(url, **k):
        raise plataforma.requests.ConnectionError("sin red")

    monkeypatch.setattr(plataforma.requests, "get", caido)
    assert plataforma.version_datos() is None


def test_montos_en_markdown_no_se_leen_como_formula():
    pytest.importorskip("streamlit")
    import sys

    sys.path.insert(0, str(RAIZ / "plataforma"))
    import comun

    # Entre dos «$» Streamlit dibuja una fórmula: el riesgo salía pegado y en cursiva.
    assert comun.md("La oferta ($881.874) supera el monto ($800.000).") == \
        r"La oferta (\$881.874) supera el monto (\$800.000)."
    assert comun.numero_texto(None) == "—" and comun.numero_texto(72450.4) == "72.450"
