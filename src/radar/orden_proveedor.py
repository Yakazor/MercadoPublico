"""Orden de compra al proveedor: la que tú le envías a la tienda o mayorista que te
vende la mercadería después de aceptar la OC del Estado.

- `crear_orden(datos)`: planilla Excel lista para imprimir o exportar a PDF. Sin datos
  es la plantilla en blanco; con datos (ver `proveedores/ordenes_proveedor/ejemplo.json`) es la
  orden llena. Subtotales, IVA y total son fórmulas.
- `validar(datos)`: lista de datos que faltan para enviar la orden.
- `totales(datos)`: subtotal, descuento, neto, IVA y total, como los calcula la planilla.
- `siguiente_numero(carpeta, anio)`: correlativo OC-AAAA-NNN según las órdenes guardadas.

El comprador puede ser persona natural con giro (nombre, RUT, giro, domicilio) o
empresa (razón social, RUT, dirección): `comprador.tipo` = "persona_natural" | "empresa".
"""
from __future__ import annotations

import json
import math
import re
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from .almacenamiento import texto_en_celda, texto_excel

IVA = 0.19
LINEAS_MINIMAS = 15
UNIDADES = ["Unidad", "Caja", "Paquete", "Resma", "Par", "Kit", "Metro", "Kilo", "Litro", "Global"]

_FUENTE = "Arial"
_ENTRADA = PatternFill("solid", fgColor="FFFF99")
_CABECERA = PatternFill("solid", fgColor="1F3864")
_BANDA = PatternFill("solid", fgColor="D9E1F2")
_LINEA = Side(style="thin", color="999999")
_CAJA = Border(left=_LINEA, right=_LINEA, top=_LINEA, bottom=_LINEA)
CLP = '$#,##0;($#,##0);"-"'

# Etiqueta y campo del comprador según su tipo, en el orden en que se muestran.
CAMPOS_COMPRADOR = {
    "persona_natural": [("Nombre completo", "nombre"), ("RUT", "rut"), ("Giro", "giro"),
                        ("Contacto", "contacto"), ("Domicilio", "domicilio")],
    "empresa": [("Razón social", "razon_social"), ("RUT", "rut"), ("Dirección", "direccion"),
                ("Contacto", "contacto")],
}
_TITULO_COMPRADOR = {"persona_natural": "COMPRADOR (persona natural con giro)",
                     "empresa": "COMPRADOR (empresa)"}
CONDICIONES = [("Lugar de entrega", "lugar_entrega"), ("Fecha límite de entrega", "fecha_entrega"),
                ("Flete", "flete"), ("Forma de pago", "forma_pago"), ("Validez y stock", "validez"),
                ("Observaciones", "observaciones")]


def _fuente(**k) -> Font:
    return Font(name=_FUENTE, **k)


def _fecha(valor):
    if isinstance(valor, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", valor):
        return date.fromisoformat(valor)
    return valor


def precio_neto(item: dict, iva: float = IVA) -> float | None:
    """Precio unitario neto del ítem: `precio_unitario_neto`, o `precio_con_iva` sin IVA
    redondeado al peso."""
    if item.get("precio_unitario_neto") not in (None, ""):
        return float(item["precio_unitario_neto"])
    if item.get("precio_con_iva") not in (None, ""):
        return float(round(float(item["precio_con_iva"]) / (1 + iva)))
    return None


def _redondeo_excel(valor: float) -> int:
    """Como ROUND(x; 0) de Excel: la mitad se aleja del cero (Python redondea al par)."""
    return int(math.floor(abs(valor) + 0.5)) * (1 if valor >= 0 else -1)


def totales(datos: dict) -> dict:
    """Los totales que calculan las fórmulas de la planilla (subtotal, descuento, neto, IVA, total)."""
    tasa = float(datos.get("iva", IVA))
    subtotal = sum(float(it["cantidad"]) * precio_neto(it, tasa) for it in datos.get("items") or []
                   if it.get("cantidad") not in (None, "") and precio_neto(it, tasa) is not None)
    neto = subtotal - float(datos.get("descuento") or 0)
    iva = _redondeo_excel(neto * tasa)
    return {"subtotal": subtotal, "descuento": float(datos.get("descuento") or 0), "neto": neto, "iva": iva,
            "total": neto + iva}


def validar(datos: dict) -> list[str]:
    """Datos que faltan para enviar la orden (vacía = completa)."""
    faltan = []
    tipo = (datos.get("comprador") or {}).get("tipo", "persona_natural")
    if tipo not in CAMPOS_COMPRADOR:
        faltan.append(f"comprador.tipo: «{tipo}» no es persona_natural ni empresa")
        tipo = "persona_natural"
    for campo in ("numero", "fecha"):
        if not datos.get(campo):
            faltan.append(campo)
    for _, campo in CAMPOS_COMPRADOR[tipo]:
        if not (datos.get("comprador") or {}).get(campo):
            faltan.append(f"comprador.{campo}")
    for campo in ("razon_social", "rut"):
        if not (datos.get("proveedor") or {}).get(campo):
            faltan.append(f"proveedor.{campo}")
    items = datos.get("items") or []
    if not items:
        faltan.append("items")
    for n, item in enumerate(items, 1):
        if not item.get("descripcion"):
            faltan.append(f"items[{n}].descripcion")
        if item.get("cantidad") in (None, ""):
            faltan.append(f"items[{n}].cantidad")
        if precio_neto(item) is None:
            faltan.append(f"items[{n}].precio")
    return faltan


def siguiente_numero(carpeta: str | Path, anio: int | None = None) -> str:
    """OC-AAAA-NNN siguiente a las órdenes guardadas en `carpeta` (*.json, sin el ejemplo)."""
    anio = anio or date.today().year
    usados = [0]
    for archivo in Path(carpeta).glob("*.json"):
        if archivo.name == "ejemplo.json":
            continue
        try:
            numero = json.loads(archivo.read_text(encoding="utf-8")).get("numero") or ""
        except (OSError, ValueError):
            continue
        m = re.fullmatch(rf"OC-{anio}-(\d+)", str(numero))
        if m:
            usados.append(int(m.group(1)))
    return f"OC-{anio}-{max(usados) + 1:03d}"


def _sin_control(valor):
    """Quita los caracteres de control que Excel no acepta (vienen en algunas descripciones)."""
    if isinstance(valor, str):
        return texto_excel(valor)
    if isinstance(valor, dict):
        return {k: _sin_control(v) for k, v in valor.items()}
    if isinstance(valor, list):
        return [_sin_control(v) for v in valor]
    return valor


def crear_orden(datos: dict | None = None) -> Workbook:
    """Planilla de la orden de compra; sin `datos`, la plantilla en blanco."""
    datos = _sin_control(datos or {})
    comprador = datos.get("comprador") or {}
    tipo = comprador.get("tipo", "persona_natural")
    if tipo not in CAMPOS_COMPRADOR:
        raise ValueError(f"comprador.tipo debe ser persona_natural o empresa, no «{tipo}»")
    proveedor = datos.get("proveedor") or {}
    referencia = datos.get("referencia") or {}
    condiciones = datos.get("condiciones") or {}
    items = datos.get("items") or []
    iva = float(datos.get("iva", IVA))

    wb = Workbook()
    ws = wb.active
    ws.title = "Orden de Compra"
    for col, ancho in zip("ABCDEFG", [6, 17, 44, 13, 10, 16, 18]):
        ws.column_dimensions[col].width = ancho

    def etiqueta(celda, texto):
        ws[celda] = texto
        ws[celda].font = _fuente(bold=True, size=10)

    def entrada(rango, valor=None, formato=None):
        if ":" in rango:
            ws.merge_cells(rango)
        valor = _fecha(valor) if valor not in (None, "") else None
        ws[rango.split(":")[0]] = valor
        texto_en_celda(ws[rango.split(":")[0]])
        if isinstance(valor, date):
            formato = "dd-mm-yyyy"
        for fila in (ws[rango] if ":" in rango else ((ws[rango],),)):
            for c in fila:
                c.fill, c.border = _ENTRADA, _CAJA
                c.font = _fuente(size=10, color="0000FF")
                c.alignment = Alignment(vertical="center", wrap_text=True)
                if formato:
                    c.number_format = formato

    def seccion(fila, texto):
        ws.merge_cells(f"A{fila}:G{fila}")
        ws[f"A{fila}"] = texto
        ws[f"A{fila}"].font = _fuente(bold=True, size=10)
        ws[f"A{fila}"].fill = _BANDA

    # Título, número y fecha
    ws.merge_cells("A1:G1")
    ws["A1"] = "ORDEN DE COMPRA"
    ws["A1"].font = _fuente(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = _CABECERA
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    etiqueta("E2", "N° OC:")
    entrada("F2:G2", datos.get("numero"))
    etiqueta("E3", "Fecha:")
    entrada("F3:G3", datos.get("fecha"), "dd-mm-yyyy")

    # Comprador: pares de campos por fila (izquierda C, derecha E:G)
    fila = 5
    seccion(fila, _TITULO_COMPRADOR[tipo])
    campos = CAMPOS_COMPRADOR[tipo]
    for i in range(0, len(campos), 2):
        fila += 1
        izq = campos[i]
        etiqueta(f"A{fila}", izq[0])
        entrada(f"C{fila}", comprador.get(izq[1]))
        if i + 1 < len(campos):
            der = campos[i + 1]
            etiqueta(f"D{fila}", der[0])
            entrada(f"E{fila}:G{fila}", comprador.get(der[1]))

    fila += 2
    seccion(fila, "PROVEEDOR (a quien le compras)")
    etiqueta(f"A{fila + 1}", "Razón social")
    entrada(f"C{fila + 1}", proveedor.get("razon_social"))
    etiqueta(f"D{fila + 1}", "RUT")
    entrada(f"E{fila + 1}:G{fila + 1}", proveedor.get("rut"))
    etiqueta(f"A{fila + 2}", "Vendedor")
    entrada(f"C{fila + 2}", proveedor.get("vendedor"))
    etiqueta(f"D{fila + 2}", "Contacto")
    entrada(f"E{fila + 2}:G{fila + 2}", proveedor.get("contacto"))

    fila += 4
    seccion(fila, "REFERENCIA MERCADO PÚBLICO (control interno)")
    etiqueta(f"A{fila + 1}", "Licitación")
    entrada(f"C{fila + 1}", referencia.get("licitacion"))
    etiqueta(f"D{fila + 1}", "ID OC Estado")
    entrada(f"E{fila + 1}:G{fila + 1}", referencia.get("oc_estado"))
    ws[f"E{fila + 1}"].comment = Comment(
        "ID de la orden de compra que te envió el organismo en mercadopublico.cl. "
        "Úsalo también en la referencia de la factura.", "Plantilla")

    # Detalle
    cabecera = fila + 3
    titulos = ["N°", "Código / SKU", "Descripción (copiar especificación exacta de las bases)",
               "Unidad", "Cantidad", "Precio unit. neto ($)", "Subtotal neto ($)"]
    for col, texto in zip("ABCDEFG", titulos):
        c = ws[f"{col}{cabecera}"]
        c.value = texto
        c.font = _fuente(bold=True, size=10, color="FFFFFF")
        c.fill, c.border = _CABECERA, _CAJA
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[cabecera].height = 30
    primera = cabecera + 1
    ultima = primera + max(LINEAS_MINIMAS, len(items)) - 1
    for r in range(primera, ultima + 1):
        n = r - primera
        ws[f"A{r}"] = n + 1
        ws[f"A{r}"].font = _fuente(size=10)
        ws[f"A{r}"].alignment = Alignment(horizontal="center")
        ws[f"A{r}"].border = _CAJA
        for col in "BCDEF":
            c = ws[f"{col}{r}"]
            c.fill, c.border = _ENTRADA, _CAJA
            c.font = _fuente(size=10, color="0000FF")
        if n < len(items):
            item = items[n]
            ws[f"B{r}"] = item.get("codigo")
            ws[f"C{r}"] = item.get("descripcion")
            ws[f"D{r}"] = item.get("unidad") or "Unidad"
            for col in "BCD":
                texto_en_celda(ws[f"{col}{r}"])
            ws[f"E{r}"] = item.get("cantidad")
            ws[f"F{r}"] = precio_neto(item, iva)
            if len(item.get("descripcion") or "") > 45:
                ws.row_dimensions[r].height = 26
        ws[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="center")
        ws[f"E{r}"].number_format = "#,##0"
        ws[f"F{r}"].number_format = CLP
        ws[f"G{r}"] = f'=IF(OR(E{r}="",F{r}=""),"",E{r}*F{r})'
        ws[f"G{r}"].number_format = CLP
        ws[f"G{r}"].border = _CAJA
        ws[f"G{r}"].font = _fuente(size=10)
    lista = DataValidation(type="list", formula1='"' + ",".join(UNIDADES) + '"', allow_blank=True)
    ws.add_data_validation(lista)
    lista.add(f"D{primera}:D{ultima}")

    # Totales
    t = ultima + 2
    totales = [("Subtotal neto", f"=SUM(G{primera}:G{ultima})", None),
               ("Descuento ($)", datos.get("descuento", 0) or 0, "entrada"),
               ("Neto", f"=F{t}-F{t + 1}", None),
               ("Tasa IVA", iva, "tasa"),
               ("IVA", f"=ROUND(F{t + 2}*F{t + 3},0)", None),
               ("TOTAL", f"=F{t + 2}+F{t + 4}", "total")]
    for k, (texto, valor, clase) in enumerate(totales):
        r = t + k
        ws.merge_cells(f"D{r}:E{r}")
        ws[f"D{r}"] = texto
        ws[f"D{r}"].font = _fuente(bold=True, size=11 if clase == "total" else 10)
        ws[f"D{r}"].alignment = Alignment(horizontal="right")
        ws.merge_cells(f"F{r}:G{r}")
        c = ws[f"F{r}"]
        c.value = valor
        c.number_format = "0%" if clase == "tasa" else CLP
        c.font = _fuente(bold=clase == "total", size=11 if clase == "total" else 10)
        c.border = ws[f"G{r}"].border = _CAJA
        if clase in ("entrada", "tasa"):
            c.fill = _ENTRADA
            c.font = _fuente(size=10, color="0000FF")

    # Condiciones
    c0 = t + 7
    seccion(c0, "CONDICIONES")
    for k, (texto, campo) in enumerate(CONDICIONES):
        r = c0 + 1 + k
        ws.merge_cells(f"A{r}:B{r}")
        etiqueta(f"A{r}", texto)
        entrada(f"C{r}:G{r}", condiciones.get(campo))
        ws.row_dimensions[r].height = 20

    # Firmas
    s = c0 + len(CONDICIONES) + 3
    firma = "Firma comprador (nombre y RUT)" if tipo == "persona_natural" else "Firma y timbre comprador"
    for rango, texto in ((f"A{s}:C{s}", firma), (f"E{s}:G{s}", "Aceptación proveedor (firma y fecha)")):
        ws.merge_cells(rango)
        a = ws[rango.split(":")[0]]
        a.value = texto
        a.font = _fuente(size=9, color="595959")
        a.alignment = Alignment(horizontal="center")
        for c in ws[rango][0]:
            c.border = Border(top=Side(style="thin", color="000000"))

    ws.print_area = f"A1:G{s}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True

    _instrucciones(wb, tipo)
    return wb


def _instrucciones(wb: Workbook, tipo: str) -> None:
    ws = wb.create_sheet("Instrucciones")
    ws.column_dimensions["A"].width = 100
    lineas = [("Cómo usar esta planilla", True),
              ("1. Llena solo las celdas AMARILLAS (texto azul). Las celdas blancas son fórmulas.", False),
              ("2. Precios NETOS (sin IVA) por unidad. Si la tienda muestra precio con IVA, divídelo por 1,19.", False),
              ("3. Copia en «Descripción» la especificación exacta que piden las bases (marca, medida, color).", False),
              ("4. Pon el ID de la OC del Estado en «ID OC Estado»: sirve para tu control y para la factura.", False),
              ("5. Exporta a PDF (Archivo → Exportar / Guardar como PDF) y envíalo al proveedor.", False),
              ("6. Pide al proveedor que confirme stock, precio y plazo por escrito antes de aceptar la OC del Estado.", False),
              ("", False),
              ("Leyenda", True),
              ("Amarillo + texto azul = dato que escribes tú.   Blanco + texto negro = fórmula, no lo edites.", False),
              ("Numeración: OC-AAAA-NNN (ej. OC-2026-001), correlativa.", False)]
    if tipo == "persona_natural":
        lineas += [("Comprador = tú como persona natural con inicio de actividades en el SII, con el mismo RUT "
                    "con que estás inscrito en Mercado Público.", False),
                   ("Como persona natural con giro vendes con factura afecta: el IVA 19 % se cobra y se "
                    "declara igual que una empresa.", False)]
    for i, (texto, negrita) in enumerate(lineas, 1):
        ws[f"A{i}"] = texto
        ws[f"A{i}"].font = _fuente(bold=negrita, size=12 if negrita else 10)
        ws[f"A{i}"].alignment = Alignment(wrap_text=True)
    ws["A10"].fill = _ENTRADA
